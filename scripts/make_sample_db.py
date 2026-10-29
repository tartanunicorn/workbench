import openpyxl, random, datetime as dt, collections, sys
from openpyxl.styles import Font, PatternFill
random.seed(7)
SRC, OUT = sys.argv[1], sys.argv[2]
src = openpyxl.load_workbook(SRC, data_only=True)
D = dt.date
SHIFT = 273  # Feb 2026 -> Nov 2026, same weekdays

def d(x):
    if x is None: return None
    if isinstance(x, dt.datetime): x = x.date()
    return x + dt.timedelta(days=SHIFT)

def addm(day, n):
    m = day.month - 1 + n
    return D(day.year + m // 12, m % 12 + 1, 1)

# ---------- owners ----------
owners = [
    ("RB", "Rekha Bawa", "Internal", "Tax ops", "Compliance Manager, run owner", 350, "Yes"),
    ("KS", "Kaitlyn Sohn", "Internal", "Tax ops", "Head of Tax, escalations", 50, "Yes"),
    ("KC", "Kristen Criswell", "Internal", "Tax ops", "Filings associate", 450, "Yes"),
    ("JW", "Jared Wallace", "Internal", "Tax ops", "Filings associate, notices", 450, "Yes"),
    ("WARP", "WARP", "Bot", "Automation", "Submits automated jurisdictions", 1000, "Yes"),
    ("MR", "MR (to confirm)", "Internal", "Unknown", "Secondary on international; identity unconfirmed", 100, "Yes"),
    ("EP", "Emily Paterson", "Internal", "Customer success", "Release owner", 600, "Yes"),
    ("MA", "Mariah", "Internal", "Customer success", "Release owner", 600, "Yes"),
    ("AM", "Austin Mahowald", "Internal", "Engineering", "Generation, diff investigation", 100, "Yes"),
]
rel_map = {"Emily": "EP", "Mariah": "MA", "Austin": "AM", "Rekha": "RB", "UNASSIGNED": "EP"}

# ---------- places ----------
ws = src["Jurisdictions"]
jur = {}
for r in ws.iter_rows(min_row=5, values_only=True):
    if r[0] and r[1] and r[2]:
        jur[r[1]] = dict(name=r[0], kind=r[2], prim=r[3], sec=r[4], warp=r[5], due=r[6], prep=r[7], note=r[12])
country_names = {"US": "United States", "CA": "Canada", "IN": "India", "EG": "Egypt", "NO": "Norway", "ZA": "South Africa",
                 "TR": "Turkey", "AE": "United Arab Emirates", "SG": "Singapore", "CH": "Switzerland", "UA": "Ukraine",
                 "EU": "European Union", "GB": "United Kingdom"}
currency = {"US": "USD", "CA": "CAD", "IN": "INR", "EG": "EGP", "NO": "NOK", "ZA": "ZAR", "TR": "TRY", "AE": "AED",
            "SG": "SGD", "CH": "CHF", "UA": "UAH", "EU": "EUR", "GB": "GBP"}
e2e = {"US-PA", "US-RI", "US-SC", "US-TN", "US-MD"}
partial = {"US-CT", "US-DC", "US-MA", "US-WA"}
def automation(code):
    return "WARP end to end" if code in e2e else "WARP partial" if code in partial else "Manual"
def prepay(code):
    return "Yes" if code in ("US-MA", "US-PA", "US-OH", "US-NY") else "Unknown"
def owner(x): return "" if x in (None, "UNASSIGNED") else x

countries, regions = [], []
for cc, nm in country_names.items():
    row = dict(country_code=cc, name=nm, is_union="Yes" if cc == "EU" else "No", currency=currency[cc],
               files_here="No", due_day=None, roll_forward="", automation="", prepayment_supported="",
               accelerated_supported="", default_preparer="", default_reviewer="", notes="")
    if cc in jur:  # national filing
        j = jur[cc]
        row.update(files_here="Yes", due_day=j["due"] if cc == "CA" else None, roll_forward="Next business day",
                   automation=automation(cc), prepayment_supported="Unknown", accelerated_supported="Unknown",
                   default_preparer=owner(j["prim"]), default_reviewer=owner(j["sec"]),
                   notes="" if j["due"] and cc == "CA" else "Due day varies: confirm")
        if j["prim"] == "UNASSIGNED": row["notes"] = "Seen in notes, no owner yet"
    countries.append(row)
for code, j in jur.items():
    if "-" not in code: continue
    parts = code.split("-")
    parent = "US-IL" if code == "US-IL-CHI" else ""
    level = "City" if code == "US-IL-CHI" else ("Province" if parts[0] == "CA" else "State")
    regions.append(dict(region_code=code, country_code=parts[0], parent_region=parent, name=j["name"], level=level,
                        files_here="Yes", due_day=j["due"], roll_forward="Next business day", automation=automation(code),
                        prepayment_supported=prepay(code), accelerated_supported="Unknown",
                        default_preparer=owner(j["prim"]), default_reviewer=owner(j["sec"]),
                        notes="Seen in notes, no owner yet" if j["prim"] == "UNASSIGNED" else ""))
regions.sort(key=lambda r: r["region_code"])

# ---------- customers ----------
ws = src["Customers"]
customers = []
for r in ws.iter_rows(min_row=5, values_only=True):
    if not r[0] or not str(r[0]).startswith("Customer"): continue
    n = int(r[0].split()[1])
    off = "Offboarding" in str(r[2])
    customers.append(dict(customer_id=f"C{n:03d}", name=r[0], org_id=r[1],
                          tier="Managed" if random.random() < .8 else "Self-serve",
                          status="Offboarding" if off else "Active",
                          release_owner=rel_map.get(r[3], "EP"), customer_approver=f"X{n:03d}",
                          historical_scope="Release with current" if random.random() < .7 else "Ask customer first",
                          internal_review=r[4] or "No",
                          notes="Offboarding: do not file" if off else ""))
cust_by_name = {c["name"]: c for c in customers}
for c in customers:
    owners.append((c["customer_approver"], f"{c['name']} finance contact", "External", c["name"], "Customer approver", None, "Yes"))
owners.append(("AUTH-NY", "NY Dept of Taxation", "External", "Tax authority", "Authority contact", None, "Yes"))

# ---------- registrations + returns from filings ----------
ws = src["Filings"]
filings = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] and str(r[0]).startswith("F-")]
regs, reg_key = [], {}
def tax_type(code):
    cc = code.split("-")[0]
    if code == "US-IL-CHI": return "Lease transaction"
    if cc == "US": return "Sales"
    if cc in ("CA", "IN"): return "GST"
    return "VAT"
for f in filings:
    key = (f[1], f[2])
    if key in reg_key: continue
    code = f[2]
    cc = code.split("-")[0]
    rid = f"REG-{len(regs)+1:04d}"
    reg_key[key] = rid
    start = random.choice([D(2024, 1, 1), D(2024, 7, 1), D(2025, 1, 1), D(2025, 4, 1), D(2025, 9, 1)])
    regs.append(dict(registration_id=rid, customer_id=cust_by_name[f[1]]["customer_id"], country_code=cc,
                     region_code=code if "-" in code else "", tax_type=tax_type(code),
                     account_source="Portal", account_ref=f"****{random.randint(1000,9999)}",
                     frequency="Monthly", active_from=start, active_to=None,
                     frequency_source="Portal", verified_by=random.choice(["RB", "KC", "JW"]),
                     prepayment_required="No", workflow="Standard",
                     last_portal_check=D(2026, random.randint(5, 10), random.randint(1, 28)), notes=""))
regs_by_id = {r["registration_id"]: r for r in regs}
hist = {reg_key[(f[1], f[2])] for f in filings if f[5] != "Current"}
# scenarios
def first_reg(region):
    return next(r for r in regs if r["region_code"] == region and r["registration_id"] not in hist)
tx = first_reg("US-TX"); tx_old = dict(tx)
tx_old.update(registration_id=f"REG-{len(regs)+1:04d}", frequency="Quarterly", active_from=D(2024, 1, 1), active_to=D(2026, 9, 30))
tx.update(active_from=D(2026, 10, 1), frequency_source="Notice", notes="Authority moved quarterly to monthly from Oct 2026")
regs.append(tx_old)
ny = first_reg("US-NY"); ny.update(frequency="Part-quarterly", prepayment_required="Yes", workflow="Prepayment", notes="Part-quarterly: monthly prepayments, quarterly reconciliation")
ms = first_reg("US-MS"); ms.update(tax_type="Use", notes="Account is use tax; check generated form")
ok = first_reg("US-OK")
regs.append(dict(ok, registration_id=f"REG-{len(regs)+1:04d}", tax_type="Vendor use", account_ref="****4471",
                 notes="Second account open; confirm which one files", active_from=D(2025, 9, 1)))
for r in random.sample([r for r in regs if r["region_code"] in ("US-MA", "US-PA", "US-OH")], 4):
    r.update(prepayment_required="Yes", workflow="Prepayment")
for r in random.sample([x for x in regs if x["registration_id"] not in hist], 25):
    if r["active_to"] is None and not r["notes"]:
        r.update(active_from=D(2026, 9, 1), notes="New registration")

run_id = "RUN-2026-11"
STAGE_FIX = {"6 Released to customer": "6 Released"}
stage_names = ["1 Not started", "2 Data synced", "3 Generated", "4 In review", "5 Ready to release",
               "6 Released", "7 Customer approved", "8 Submitted", "9 Confirmed / paid"]
returns, issue_links = [], []
for i, f in enumerate(filings, 1):
    rid = f"RET-{i:04d}"
    reg = reg_key[(f[1], f[2])]
    p = f[4].date() if isinstance(f[4], dt.datetime) else f[4]
    pstart = addm(p, 9)
    pend = addm(pstart, 1) - dt.timedelta(days=1)
    kind = "Regular" if f[5] == "Current" else "Historical"
    st = STAGE_FIX.get(f[6], f[6])
    state, stage = "Open", st
    if st == "Blocked": state, stage = "Blocked", random.choice(stage_names[1:4])
    elif st == "On hold": state, stage = "On hold", stage_names[0]
    elif st == "Not required": state, stage = "Not required", stage_names[0]
    elif st and st.startswith("9"): state = "Done"
    calc = round(random.lognormvariate(7.2, 1.3), 2)
    r = random.random()
    coll = round(calc * (1 + (random.uniform(-.02, .02) if r < .88 else random.uniform(.06, .3) * random.choice([-1, 1]))), 2)
    sn = int(stage[0])
    rel = d(f[12]) if f[12] else None
    appr = None
    if sn >= 7 and rel: appr = min(rel + dt.timedelta(days=random.randint(0, 2)), D(2026, 11, 5))
    sub = d(f[13]) if f[13] else None
    returns.append(dict(return_id=rid, registration_id=reg, run_id=run_id, period_start=pstart, period_end=pend,
                        kind=kind, stage=stage, state=state, calculated=calc, collected=coll, due=calc,
                        statutory_due=None, released_on=rel, approved_on=appr, submitted_on=sub,
                        confirmation_no=f"CNF-{random.randint(100000,999999)}" if sub else "",
                        reason=f[19] or ("No sales in period; zero return not required" if state == "Not required" else ""),
                        notes=""))
    if f[15]: issue_links.append(dict(issue_id=f[15], return_id=rid, registration_id="", customer_id=""))
# TX scenario: the Oct return links to the new monthly registration (already), mark note
for rt in returns:
    if rt["registration_id"] == tx["registration_id"] and rt["kind"] == "Regular":
        rt["notes"] = "Expected only because of the Oct frequency change"

# semi-monthly scenario: one registration files twice in the November run
sm = next(r for r in regs if r["region_code"] == "US-CO" and r["registration_id"] not in hist and not r["notes"])
sm.update(frequency="Semi-monthly", notes="Semi-monthly filer: two returns per month, each with its own due date")
smr = next(rt for rt in returns if rt["registration_id"] == sm["registration_id"])
smr.update(period_start=D(2026, 10, 16), period_end=D(2026, 10, 31), statutory_due=D(2026, 11, 10), notes="Second half of October")
returns.append(dict(smr, return_id=f"RET-{len(returns)+1:04d}", period_start=D(2026, 11, 1), period_end=D(2026, 11, 15),
                    statutory_due=D(2026, 11, 25), stage="1 Not started", state="Open", released_on=None, approved_on=None,
                    submitted_on=None, confirmation_no="", notes="First half of November; period still open"))
# ---------- history: the closed October run (September periods) ----------
oct_run = "RUN-2026-10"
nov_regular = [rt for rt in returns if rt["kind"] == "Regular"]
hist_returns = []
def add_hist(reg_id, ps, pe, stat=None):
    calc = round(random.lognormvariate(7.2, 1.3), 2)
    rel = D(2026, 10, random.choice([6, 7, 7, 7, 8]))
    appr = rel + dt.timedelta(days=random.randint(0, 4))
    late = random.random() < 0.03
    sub = D(2026, 10, random.randint(20, 23)) if late else D(2026, 10, random.randint(13, 19))
    hist_returns.append(dict(return_id=None, registration_id=reg_id, run_id=oct_run, period_start=ps, period_end=pe, kind="Regular",
                             stage="9 Confirmed / paid", state="Done", calculated=calc, collected=round(calc * random.uniform(.98, 1.02), 2), due=calc,
                             statutory_due=stat, released_on=rel, approved_on=appr, submitted_on=sub, confirmation_no=f"CNF-{random.randint(100000,999999)}",
                             reason="", notes="Filed after CD19" if late else ""))
seen = set()
for rt in nov_regular:
    reg = regs_by_id[rt["registration_id"]]
    if reg["registration_id"] in seen: continue
    seen.add(reg["registration_id"])
    if reg["active_from"] > D(2026, 9, 30): continue
    if reg["frequency"] == "Semi-monthly":
        add_hist(reg["registration_id"], D(2026, 9, 16), D(2026, 9, 30), D(2026, 10, 10))
        add_hist(reg["registration_id"], D(2026, 10, 1), D(2026, 10, 15), D(2026, 10, 25))
    else:
        add_hist(reg["registration_id"], D(2026, 9, 1), D(2026, 9, 30))
# TX quarterly registration files its Q3 return in the October run, then closes
add_hist(tx_old["registration_id"], D(2026, 7, 1), D(2026, 9, 30))
for h in hist_returns:
    h["return_id"] = f"RET-{len(returns)+1:04d}"
    returns.append(h)

# ---------- ownerships ----------
ownerships = []
def place_defaults(reg):
    if reg["region_code"]:
        rg = next(x for x in regions if x["region_code"] == reg["region_code"])
        return rg["default_preparer"], rg["default_reviewer"]
    c = next(x for x in countries if x["country_code"] == reg["country_code"])
    return c["default_preparer"], c["default_reviewer"]
cust_by_id = {c["customer_id"]: c for c in customers}
regs_by_id = {r["registration_id"]: r for r in regs}
for rt in returns:
    reg = regs_by_id[rt["registration_id"]]
    pp, rv = place_defaults(reg)
    cu = cust_by_id[reg["customer_id"]]
    for role, o in (("Preparer", pp), ("Reviewer", rv), ("Release owner", cu["release_owner"]), ("Customer approver", cu["customer_approver"])):
        if o: ownerships.append(dict(owner_id=o, role=role, return_id=rt["return_id"], run_id="", issue_id="", from_date=D(2026, 10, 1) if rt["run_id"] == oct_run else D(2026, 11, 1), to_date=None))
# a reassignment example
ex = next(o for o in ownerships if o["role"] == "Preparer" and o["owner_id"] == "KC")
ex["to_date"] = D(2026, 11, 3)
ownerships.append(dict(ex, owner_id="JW", from_date=D(2026, 11, 3), to_date=None))
ownerships += [dict(owner_id="RB", role="Run owner", return_id="", run_id=oct_run, issue_id="", from_date=D(2026, 10, 1), to_date=None),
               dict(owner_id="RB", role="Run owner", return_id="", run_id=run_id, issue_id="", from_date=D(2026, 11, 1), to_date=None),
               dict(owner_id="KS", role="Escalation", return_id="", run_id=run_id, issue_id="", from_date=D(2026, 11, 1), to_date=None)]

# ---------- issues / balances ----------
ws = src["Issues"]
issues, balances = [], []
anon = {}
def anon_c(name):
    if name not in anon: anon[name] = random.choice(customers[:110])["customer_id"]
    return anon[name]
for r in ws.iter_rows(min_row=5, values_only=True):
    if not r[0] or not str(r[0]).startswith("ISS"): continue
    iid, opened, cust, jur_s, cat, sev, desc, det, own, status, due_by, nxt, sysfix, link = r[:14]
    if cat == "Notice / balance due" or (own == "JW" and det == "Tax authority notice" and cat != "Account / portal access"):
        balances.append(dict(notice_id=f"BN-{len(balances)+1:03d}", customer_id=anon_c(cust), place=jur_s,
                             kind="Credit" if "Credit" in desc else "Reversed payment" if "reversed" in desc else "Balance due",
                             period=desc, amount=None, received_on=D(2026, 10, random.randint(1, 28)),
                             respond_by=D(2026, 11, random.randint(10, 30)), status="Open", owner_id="JW", issue_id=""))
        continue
    is_sample = str(cust).startswith("several")
    desc = str(desc).replace("SAMPLE: ", "").replace("Jan return", "Oct return")
    issues.append(dict(issue_id=iid, category=cat, severity=(sev or "")[:2], status=status,
                       opened_on=d(opened) if opened else D(2026, 10, random.randint(1, 30)),
                       due_by=D(2026, 11, 5) if (sev or "").startswith("S1") else None,
                       description=desc, detected_by=det or "", resolution=nxt or "",
                       systemic_fix="Yes" if sysfix == "Yes" else "No", source="Linear" if link else "",
                       source_ref=link.split("/issue/")[1].split("/")[0] if link else "",
                       customer_hint="" if is_sample else anon_c(cust), place_hint=jur_s or ""))
    if own: ownerships.append(dict(owner_id={"Austin": "AM"}.get(own, own), role="Issue owner", return_id="", run_id="", issue_id=iid, from_date=None, to_date=None))
for i in issues:
    if i["customer_hint"] and i["status"] != "Resolved":
        issue_links.append(dict(issue_id=i["issue_id"], return_id="", registration_id="", customer_id=i["customer_hint"]))
ownerships.append(dict(owner_id="X014", role="Waiting on", return_id="", run_id="", issue_id="ISS-024", from_date=None, to_date=None))

# ---------- events ----------
events = []
def ev(rid, at, fs, ts, fst, tst, actor, src_, note):
    events.append(dict(event_id=f"EV-{len(events)+1:05d}", return_id=rid, at=at, from_stage=fs, to_stage=ts,
                       from_state=fst, to_state=tst, actor=actor, source=src_, note=note))
for rt in returns[:40]:
    sn = int(rt["stage"][0])
    t = dt.datetime(2026, 11, 2, 9, 0)
    for s in range(1, sn):
        t += dt.timedelta(hours=random.randint(2, 14))
        ev(rt["return_id"], t, stage_names[s-1], stage_names[s], "Open", "Open",
           "WARP" if s == 7 else random.choice(["KC", "JW", "RB"]), "Person" if s != 7 else "WARP", "")
    if rt["state"] == "Blocked":
        ev(rt["return_id"], t + dt.timedelta(hours=1), rt["stage"], rt["stage"], "Open", "Blocked", "KC", "Person", "Blocked by linked issue")

# ---------- runs, holidays, settings ----------
runs = [dict(run_id="RUN-2026-10", run_month=D(2026, 10, 1), period_start=D(2026, 9, 1), period_end=D(2026, 9, 30), status="Closed", notes="Previous run"),
        dict(run_id=run_id, run_month=D(2026, 11, 1), period_start=D(2026, 10, 1), period_end=D(2026, 10, 31), status="Open", notes="")]
holidays = [D(2026, 1, 1), D(2026, 1, 19), D(2026, 2, 16), D(2026, 5, 25), D(2026, 6, 19), D(2026, 7, 3), D(2026, 9, 7),
            D(2026, 10, 12), D(2026, 11, 11), D(2026, 11, 26), D(2026, 12, 25), D(2027, 1, 1), D(2027, 1, 18)]
settings = [("as_of_date", D(2026, 11, 5), "Fixed to Thu 5 Nov 2026 (BD4) for the demo. Leave blank to use today."),
            ("current_run", run_id, "Run shown by default"),
            ("release_business_day", 5, "Release to customers by this business day"),
            ("submit_calendar_day", 19, "Submit to authority by this calendar day"),
            ("submit_buffer_days", 2, "Internal target is this many days before the submit day"),
            ("review_gap_amount", 1000, "Second review if calculated vs collected gap exceeds this amount"),
            ("review_gap_pct", 0.05, "...or this share of calculated"),
            ("new_registration_days", 60, "Second review if registration started within this many days of period end"),
            ("at_risk_days", 3, "Submit at risk when this many days or fewer remain")]

# ---------- tasks ----------
hol = set(holidays)
def isbiz(x): return x.weekday() < 5 and x not in hol
def bd(month_start, n):
    x = month_start
    if n > 0:
        c = 0
        while True:
            if isbiz(x):
                c += 1
                if c == n: return x
            x += dt.timedelta(days=1)
    c = 0; x = month_start - dt.timedelta(days=1)
    while True:
        if isbiz(x):
            c += 1
            if c == -n + 1: return x
        x -= dt.timedelta(days=1)
def due_for(anchor, off, ms):
    return bd(ms, off) if anchor == "BD" else D(ms.year, ms.month, off)
templates = [
    ("Pre-run check", "Check frequency changes from notices and portals", "BD", -2, "RB"),
    ("Pre-run check", "Confirm new registrations and closed accounts", "BD", -2, "KC"),
    ("Pre-run check", "Confirm offboarding and on-hold customers", "BD", -1, "EP"),
    ("Pre-run check", "List prepayments due this month", "BD", -1, "JW"),
    ("Pre-run check", "Every filing place has an owner", "BD", 0, "RB"),
    ("Milestone", "Data synced; sync errors and missing PTCs logged as issues", "BD", 1, "AM"),
    ("Milestone", "All current returns generated or blocked with an issue", "BD", 2, "RB"),
    ("Milestone", "Go / no-go call: what releases tomorrow", "BD", 4, "RB"),
    ("Milestone", "Everything released or customer told why", "BD", 5, "RB"),
    ("Chase", "First chase on unapproved returns", "CD", 8, "EP"),
    ("Chase", "Second chase on unapproved returns", "CD", 11, "MA"),
    ("Chase", "Escalate unapproved returns to CS lead", "CD", 12, "RB"),
    ("Milestone", "Anything not submitted goes to Rekha", "CD", 17, "RB"),
    ("Milestone", "Close-out: metrics and 30-minute retro", "CD", 22, "KS"),
]
# which returns each run task is about, as a filter the app opens on click
RFILTER = {"Confirm offboarding and on-hold customers": "state=On hold", "List prepayments due this month": "prepay=Yes",
           "Every filing place has an owner": "unowned=Yes", "Data synced; sync errors and missing PTCs logged as issues": "stage<2; state=Open",
           "All current returns generated or blocked with an issue": "stage<3; state=Open", "Go / no-go call: what releases tomorrow": "stage=5 Ready to release",
           "Everything released or customer told why": "stage<6; state=Open", "First chase on unapproved returns": "stage=6 Released",
           "Second chase on unapproved returns": "stage=6 Released", "Escalate unapproved returns to CS lead": "stage=6 Released",
           "Anything not submitted goes to Rekha": "stage<8; state=Open", "Check frequency changes from notices and portals": "flag=Attention",
           "Confirm new registrations and closed accounts": "tier=A", "Close-out: metrics and 30-minute retro": "flag=Risk*"}
tasks = []
as_of = D(2026, 11, 5)
ms = D(2026, 11, 1)
for i, (kind, title, anc, off, who) in enumerate(templates, 1):
    tid = f"TSK-{i:03d}"
    tasks.append(dict(task_id=tid, title=title, kind=kind, run_id=None, template_id=None, return_id=None, issue_id=None,
                      customer_id=None, notice_id=None, anchor=anc, offset=off, due_date=None, status="Template", done_on=None, notes="Template: copied into each run",
                      return_filter=RFILTER.get(title)))
    ownerships.append(dict(owner_id=who, role="Task owner", return_id="", run_id="", issue_id="", task_id=tid, from_date=None, to_date=None))
for t in list(tasks):
    n = len(tasks) + 1
    due = due_for(t["anchor"], t["offset"], ms)
    status = "Done" if due < as_of else ("In progress" if due == as_of else "Not started")
    tid = f"TSK-{n:03d}"
    tasks.append(dict(t, task_id=tid, run_id=run_id, template_id=t["task_id"], due_date=due, status=status,
                      done_on=due if status == "Done" else None, notes=None))
    who = next(o["owner_id"] for o in ownerships if o.get("task_id") == t["task_id"])
    ownerships.append(dict(owner_id=who, role="Task owner", return_id="", run_id="", issue_id="", task_id=tid, from_date=D(2026, 11, 1), to_date=None))
for t in [x for x in tasks if x["run_id"] is None]:
    n = len(tasks) + 1
    due = due_for(t["anchor"], t["offset"], D(2026, 10, 1))
    tasks.append(dict(t, task_id=f"TSK-{n:03d}", run_id=oct_run, template_id=t["task_id"], due_date=due, status="Done", done_on=due, notes=None))
    who = next(o["owner_id"] for o in ownerships if o.get("task_id") == t["task_id"])
    ownerships.append(dict(owner_id=who, role="Task owner", return_id="", run_id="", issue_id="", task_id=f"TSK-{n:03d}", from_date=D(2026, 10, 1), to_date=None))
# the pre-run owner check is overdue because 6 places have no owner
for t in tasks:
    if t["title"].startswith("Every filing place") and t["run_id"] == run_id: t.update(status="In progress", done_on=None, notes="NJ, SG, CH, UA, EU, GB still have no owner")
for b in balances:
    n = len(tasks) + 1
    tasks.append(dict(task_id=f"TSK-{n:03d}", title=f"Respond to {b['kind'].lower()}: {b['period'][:50]}", kind="Notice response", run_id=run_id, template_id=None,
                      return_id=None, issue_id=None, customer_id=None, notice_id=b["notice_id"], anchor=None, offset=None, due_date=b["respond_by"],
                      status="Not started", done_on=None, notes=None))
    ownerships.append(dict(owner_id="JW", role="Task owner", return_id="", run_id="", issue_id="", task_id=f"TSK-{n:03d}", from_date=D(2026, 11, 1), to_date=None))
for iss in issues:
    if iss["severity"] == "S1":
        n = len(tasks) + 1
        tasks.append(dict(task_id=f"TSK-{n:03d}", title=f"Follow up {iss['issue_id']} before release", kind="Follow-up", run_id=run_id, template_id=None,
                          return_id=None, issue_id=iss["issue_id"], customer_id=None, notice_id=None, anchor=None, offset=None, due_date=D(2026, 11, 5),
                          status="In progress", done_on=None, notes=None))
        ownerships.append(dict(owner_id="RB", role="Task owner", return_id="", run_id="", issue_id="", task_id=f"TSK-{n:03d}", from_date=D(2026, 11, 1), to_date=None))
n = len(tasks) + 1
off_c = next(c for c in customers if c["status"] == "Offboarding")
tasks.append(dict(task_id=f"TSK-{n:03d}", title="Confirm final return and close registrations", kind="Follow-up", run_id=run_id, template_id=None,
                  return_id=None, issue_id=None, customer_id=off_c["customer_id"], notice_id=None, anchor=None, offset=None, due_date=D(2026, 11, 12),
                  status="Not started", done_on=None, notes="Customer is offboarding"))
ownerships.append(dict(owner_id="EP", role="Task owner", return_id="", run_id="", issue_id="", task_id=f"TSK-{n:03d}", from_date=D(2026, 11, 1), to_date=None))
for o in ownerships: o.setdefault("task_id", "")

# ---------- return types: the master list each registration points at ----------
place_by = {c["country_code"]: c for c in countries}
place_by.update({r["region_code"]: r for r in regions})
cust_by_id = {c["customer_id"]: c for c in customers}
rtypes, rt_key = [], {}
def type_for(reg):
    code = reg["region_code"] or reg["country_code"]
    pl = place_by[code]
    wf = "WARP" if pl["automation"] == "WARP end to end" else ("Prepayment" if reg["prepayment_required"] == "Yes" else "Standard")
    key = (reg["country_code"], reg["region_code"], reg["tax_type"], reg["frequency"], wf)
    if key not in rt_key:
        rid = f"RT-{len(rtypes)+1:03d}"
        semi = reg["frequency"] in ("Semi-monthly", "Weekly")
        rtypes.append(dict(return_type_id=rid, name=f"{code} {reg['tax_type']} · {reg['frequency']}{' · '+wf if wf!='Standard' else ''}",
                           country_code=reg["country_code"], region_code=reg["region_code"] or None, tax_type=reg["tax_type"],
                           frequency=reg["frequency"], workflow=wf,
                           due_rule="Days after period end" if semi else "Day of next month",
                           due_value=10 if semi else (pl["due_day"] or None),
                           roll_forward="Next business day", automation=pl["automation"] or "Manual",
                           default_preparer=pl["default_preparer"] or None, default_reviewer=pl["default_reviewer"] or None,
                           active="Yes", notes=None if pl["due_day"] or semi else "Due day varies: confirm"))
        rt_key[key] = rid
    return rt_key[key]
for reg in regs:
    reg["return_type_id"] = type_for(reg)
# places seen only in notes: a return type with no owner yet
for code, tt in (("US-NJ", "Sales"), ("SG", "GST"), ("CH", "VAT"), ("UA", "VAT"), ("EU", "VAT"), ("GB", "VAT")):
    cc = code.split("-")[0]
    rtypes.append(dict(return_type_id=f"RT-{len(rtypes)+1:03d}", name=f"{code} {tt} · Monthly", country_code=cc, region_code=code if "-" in code else None,
                       tax_type=tt, frequency="Monthly", workflow="Standard", due_rule="Day of next month", due_value=None,
                       roll_forward="Next business day", automation="Manual", default_preparer=None, default_reviewer=None,
                       active="Yes", notes="Seen in tracker notes; no owner or due day yet"))
rt_by = {t["return_type_id"]: t for t in rtypes}

# rules that change over time live in ReturnTypeRules: one row per field change, valid from a date.
# The value for a period is the latest row on or before its period end, field by field.
RULE_FIELDS = ["due_rule", "due_value", "roll_forward", "automation", "default_preparer", "default_reviewer"]
rules = []
START = D(2026, 1, 1)
for t in rtypes:
    for f in RULE_FIELDS:
        if t.get(f) not in (None, ""): rules.append(dict(return_type_id=t["return_type_id"], field=f, value=t[f], valid_from=START, notes=None))
# sample history and stacked future changes, to show that a later change never drops an earlier one
def first_type(pred): return next(t for t in rtypes if pred(t))
tx = first_type(lambda t: t["region_code"] == "US-TX" and t["frequency"] == "Monthly")
ca = first_type(lambda t: t["region_code"] == "US-CA")
nj = first_type(lambda t: t["region_code"] == "US-NJ")
rules += [
    dict(return_type_id=ca["return_type_id"], field="automation", value="WARP partial", valid_from=D(2026, 9, 1), notes="WARP now generates CA; we still submit"),
    dict(return_type_id=tx["return_type_id"], field="due_value", value=25, valid_from=D(2026, 12, 1), notes="Change 1: example due day change"),
    dict(return_type_id=tx["return_type_id"], field="automation", value="WARP end to end", valid_from=D(2027, 1, 1), notes="Change 2: WARP takes over; due day from change 1 still applies"),
    dict(return_type_id=nj["return_type_id"], field="due_value", value=20, valid_from=START, notes="NJ ST-50/51 due the 20th; owner still needed"),
]
for i, r in enumerate(rules, 1): r["rule_id"] = f"RTR-{i:05d}"
def rule_at(type_id, field, on):
    best = None
    for r in rules:
        if r["return_type_id"] == type_id and r["field"] == field and r["valid_from"] <= on and (best is None or r["valid_from"] >= best["valid_from"]): best = r
    return best["value"] if best else None
def type_at(type_id, on):
    t = dict(rt_by[type_id])
    for f in RULE_FIELDS: t[f] = rule_at(type_id, f, on)
    return t

# statutory due on every return, from its return type
def stat_due(rt, run_month, pend):
    if rt["due_rule"] == "Days after period end": return pend + dt.timedelta(days=int(rt["due_value"]))
    return D(run_month.year, run_month.month, int(rt["due_value"])) if rt["due_value"] else None
run_month_of = {"RUN-2026-10": D(2026, 10, 1), "RUN-2026-11": D(2026, 11, 1)}
for rt_ in returns:
    t = type_at(regs_by_id[rt_["registration_id"]]["return_type_id"], rt_["period_end"])
    if not rt_.get("statutory_due"): rt_["statutory_due"] = stat_due(t, run_month_of[rt_["run_id"]], rt_["period_end"])

# checklist templates per return type
STEPS = {
 "Standard": [("Sync data and clear blockers","BD",1,2,"Preparer","All"),("Generate return","BD",2,3,"Preparer","All"),
              ("Preparer checks: calc vs collected, credits, frequency","BD",3,4,"Preparer","All"),("Second review","BD",4,5,"Reviewer","A"),
              ("Release to customer with note","BD",5,6,"Release owner","All"),("Customer approval (chase CD8, CD11)","CD",12,7,"Release owner","All"),
              ("Submit and save confirmation","DUE",-2,8,"Preparer","All"),("Confirm payment debited","DUE",3,9,"Preparer","All")],
}
STEPS["WARP"] = [s if s[0] != "Submit and save confirmation" else ("WARP submits; check confirmation and PDF","DUE",-2,8,"Preparer","All") for s in STEPS["Standard"]]
STEPS["Prepayment"] = STEPS["Standard"][:2] + [("Calculate and submit prepayment","BD",3,4,"Preparer","All")] + STEPS["Standard"][2:]
for t in tasks:
    t.update(return_type_id=None, completes_stage=None, applies_to_tier=None, owner_role=None)
tpl_by_type = {}
for t in rtypes:
    for (title, anc, off, stg, role, tier) in STEPS[t["workflow"]]:
        n = len(tasks) + 1
        row_ = dict(task_id=f"TSK-{n:05d}", title=title, kind="Return step", run_id=None, template_id=None, return_type_id=t["return_type_id"],
                    return_id=None, issue_id=None, customer_id=None, notice_id=None, anchor=anc, offset=off, completes_stage=stg,
                    applies_to_tier=tier, owner_role=role, due_date=None, status="Template", done_on=None, notes=None)
        tasks.append(row_); tpl_by_type.setdefault(t["return_type_id"], []).append(row_)
# tier, same rules as the app
def tier(rt_):
    reg = regs_by_id[rt_["registration_id"]]; cu = cust_by_id[reg["customer_id"]]
    calc = rt_["calculated"] or 0; gap = abs((rt_["collected"] or 0) - calc)
    if rt_["kind"] != "Regular" or gap > 1000 or (calc and gap / calc > .05) or (rt_["period_end"] - reg["active_from"]).days <= 60 \
       or reg["prepayment_required"] == "Yes" or cu["internal_review"] == "Yes" or any(l["return_id"] == rt_["return_id"] for l in issue_links):
        return "A"
    return "B"
def step_due(anc, off, run_month, sub_due):
    if anc == "DUE": return sub_due + dt.timedelta(days=off)
    return due_for(anc, off, run_month)
for rt_ in returns:
    t = rt_by[regs_by_id[rt_["registration_id"]]["return_type_id"]]
    rm = run_month_of[rt_["run_id"]]
    cd19 = D(rm.year, rm.month, 19)
    sub_due = min(cd19, rt_["statutory_due"]) if rt_["statutory_due"] else cd19
    sn = int(rt_["stage"][0]); tr = tier(rt_)
    nxt = True
    for tp in tpl_by_type[t["return_type_id"]]:
        if tp["applies_to_tier"] == "A" and tr != "A": continue
        due = step_due(tp["anchor"], tp["offset"], rm, sub_due)
        if tp["completes_stage"] <= sn or rt_["state"] == "Done": st, done = "Done", min(due, as_of)
        elif rt_["state"] in ("Not required",): st, done = "Skipped", None
        elif nxt and rt_["state"] in ("Open", "Blocked"): st, done, nxt = "In progress", None, False
        else: st, done = "Not started", None
        n = len(tasks) + 1
        tasks.append(dict(task_id=f"TSK-{n:05d}", title=tp["title"], kind="Return step", run_id=rt_["run_id"], template_id=tp["task_id"],
                          return_type_id=None, return_id=rt_["return_id"], issue_id=None, customer_id=None, notice_id=None,
                          anchor=tp["anchor"], offset=tp["offset"], completes_stage=tp["completes_stage"], applies_to_tier=tp["applies_to_tier"],
                          owner_role=tp["owner_role"], due_date=due, status=st, done_on=done, notes=None))
# places become geography only
for c in countries:
    for k in ("files_here","due_day","roll_forward","automation","prepayment_supported","accelerated_supported","default_preparer","default_reviewer"): c.pop(k, None)
for r in regions:
    for k in ("files_here","due_day","roll_forward","automation","prepayment_supported","accelerated_supported","default_preparer","default_reviewer"): r.pop(k, None)
    if r["notes"] and "owner" in r["notes"]: r["notes"] = None
for c in countries:
    if c.get("notes") and ("owner" in c["notes"] or "Due day" in c["notes"]): c["notes"] = None

# header tasks: every run and every return is also a task, so the whole job is one graph
rels = []
def rel(a, b, typ, note=None): rels.append(dict(from_task_id=a, to_task_id=b, relation_type=typ, notes=note))
run_head, ret_head = {}, {}
for r in runs:
    n = len(tasks) + 1; tid = f"TSK-{n:05d}"
    tasks.append(dict(task_id=tid, title=f"{r['run_month']:%b %Y} filing run", kind="Run", run_id=r["run_id"], template_id=None, return_type_id=None,
                      return_id=None, issue_id=None, customer_id=None, notice_id=None, anchor=None, offset=None, completes_stage=None,
                      applies_to_tier=None, owner_role=None, due_date=None, status=None, done_on=None, notes="Header: status and dates follow the Runs row"))
    run_head[r["run_id"]] = tid
for rt_ in returns:
    t = rt_by[regs_by_id[rt_["registration_id"]]["return_type_id"]]
    n = len(tasks) + 1; tid = f"TSK-{n:05d}"
    tasks.append(dict(task_id=tid, title=f"{t['name']} · {rt_['period_end']:%b %Y}" + ("" if rt_["kind"] == "Regular" else f" · {rt_['kind']}"),
                      kind="Return", run_id=rt_["run_id"], template_id=None, return_type_id=None, return_id=rt_["return_id"], issue_id=None,
                      customer_id=None, notice_id=None, anchor=None, offset=None, completes_stage=None, applies_to_tier=None, owner_role=None,
                      due_date=None, status=None, done_on=None, notes="Header: status and dates follow the Returns row"))
    ret_head[rt_["return_id"]] = tid
    rel(tid, run_head[rt_["run_id"]], "Subtask of")
steps_of = {}
for t in tasks:
    if t["kind"] == "Return step" and t["return_id"]:
        rel(t["task_id"], ret_head[t["return_id"]], "Subtask of")
        steps_of.setdefault(t["return_id"], []).append(t)
        t["return_id"] = None
    elif t["run_id"] and t["kind"] not in ("Run", "Return"):
        rel(t["task_id"], run_head[t["run_id"]], "Subtask of")
# an S1 issue follow-up blocks release of every return it touches
for t in tasks:
    if t["kind"] == "Follow-up" and t["issue_id"]:
        for l in issue_links:
            if l["issue_id"] == t["issue_id"] and l["return_id"]:
                st = next((x for x in steps_of.get(l["return_id"], []) if x["completes_stage"] == 6 and x["status"] not in ("Done", "Skipped")), None)
                if st: rel(t["task_id"], st["task_id"], "Blocks")
# read order: templates, then each run: header, run tasks, then each return header followed by its steps
pos = {}
order = [t for t in tasks if t["run_id"] is None]
by_run_misc = {}
for t in tasks:
    if t["run_id"] and t["kind"] not in ("Run", "Return", "Return step"): by_run_misc.setdefault(t["run_id"], []).append(t)
tby = {t["task_id"]: t for t in tasks}
for r in runs:
    order.append(tby[run_head[r["run_id"]]]); order += by_run_misc.get(r["run_id"], [])
    for rt_ in returns:
        if rt_["run_id"] == r["run_id"]:
            order.append(tby[ret_head[rt_["return_id"]]]); order += steps_of.get(rt_["return_id"], [])
assert len(order) == len(tasks), (len(order), len(tasks))
tasks = order

# one id format for every task
remap = {t["task_id"]: f"TSK-{i:05d}" for i, t in enumerate(tasks, 1)}
for t in tasks:
    t["task_id"] = remap[t["task_id"]]
    if t["template_id"]: t["template_id"] = remap[t["template_id"]]
for o in ownerships:
    if o.get("task_id"): o["task_id"] = remap[o["task_id"]]
for i, r in enumerate(rels, 1):
    r["from_task_id"], r["to_task_id"] = remap[r["from_task_id"]], remap[r["to_task_id"]]
    r["relation_id"] = f"REL-{i:06d}"

# ---------- write ----------
wb = openpyxl.Workbook(); wb.remove(wb.active)
HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F3A5F")
def sheet(name, rows, cols=None, note=None):
    ws = wb.create_sheet(name)
    cols = cols or list(rows[0].keys())
    ws.append(cols)
    for c in ws[1]: c.font, c.fill = HEAD, FILL
    for r in rows: ws.append([r.get(c) if isinstance(r, dict) else r[i] for i, c in enumerate(cols)])
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, 1):
        w = max([len(str(c))] + [len(str(ws.cell(j, i).value or "")) for j in range(2, min(ws.max_row, 200) + 1)])
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(10, w + 2), 60)
        for j in range(2, ws.max_row + 1):
            v = ws.cell(j, i).value
            if isinstance(v, dt.datetime): ws.cell(j, i).number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(v, dt.date): ws.cell(j, i).number_format = "yyyy-mm-dd"
    return ws
sheet("Settings", settings, ["key", "value", "notes"])
sheet("Runs", runs)
sheet("Owners", owners, ["owner_id", "name", "type", "team_or_company", "role_note", "capacity_per_run", "active"])
sheet("Ownerships", ownerships, ["owner_id", "role", "return_id", "run_id", "issue_id", "task_id", "from_date", "to_date"])
sheet("Countries", countries)
sheet("Regions", regions)
sheet("Customers", customers)
sheet("ReturnTypes", [{k: v for k, v in t.items() if k not in RULE_FIELDS} for t in rtypes])
sheet("ReturnTypeRules", rules, ["rule_id","return_type_id","field","value","valid_from","notes"])
sheet("Registrations", regs, ["registration_id","customer_id","return_type_id","account_source","account_ref","active_from","active_to","frequency_source","verified_by","last_portal_check","notes"])
sheet("Returns", returns)
sheet("ReturnEvents", events)
issues_out = [{k: v for k, v in i.items() if k not in ("customer_hint", "place_hint")} for i in issues]
sheet("Issues", issues_out)
sheet("IssueLinks", issue_links)
sheet("BalancesNotices", balances)
sheet("Tasks", tasks, ["task_id","title","kind","return_filter","run_id","template_id","return_type_id","return_id","issue_id","customer_id","notice_id","anchor","offset","completes_stage","applies_to_tier","owner_role","due_date","status","done_on","notes"])
sheet("TaskRelations", rels, ["relation_id","from_task_id","to_task_id","relation_type","notes"])
sheet("Holidays", [(h, "US federal") for h in holidays], ["date", "calendar"])
wb.save(OUT)
print(len(owners), len(ownerships), len(countries), len(regions), len(customers), len(regs), len(returns), len(events), len(issues), len(issue_links), len(balances), len(tasks), len(rels))
